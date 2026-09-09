from io import BytesIO

import fitz
import pytest
from docx import Document

from app.core.exceptions import AppException
from app.services.text_extraction_service import TextExtractionService


def create_pdf(path, text: str) -> None:
    document = fitz.open()
    document.new_page().insert_text((72, 72), text)
    document.save(path)
    document.close()


def create_docx(path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "FastAPI"
    document.save(path)


def test_extracts_pdf_text_and_metadata(tmp_path) -> None:
    path = tmp_path / "resume.pdf"
    create_pdf(path, "Jordan Doe\nPython Developer")

    result = TextExtractionService().extract_text(path)

    assert result.text == "Jordan Doe\nPython Developer"
    assert result.page_count == 1
    assert result.character_count == len(result.text)
    assert result.as_dict()["text"] == result.text


def test_extracts_docx_paragraphs_and_tables(tmp_path) -> None:
    path = tmp_path / "resume.docx"
    create_docx(path, "Jordan Doe")

    result = TextExtractionService().extract_text(path, "docx")

    assert result.text == "Jordan Doe\nPython | FastAPI"
    assert result.page_count == 1


@pytest.mark.parametrize("suffix", ["pdf", "docx"])
def test_rejects_document_without_extractable_text(tmp_path, suffix: str) -> None:
    path = tmp_path / f"empty.{suffix}"
    if suffix == "pdf":
        document = fitz.open()
        document.new_page()
        document.save(path)
        document.close()
    else:
        Document().save(path)

    with pytest.raises(AppException) as error:
        TextExtractionService().extract_text(path)

    assert error.value.code == "NO_EXTRACTABLE_TEXT"


def test_rejects_malformed_document(tmp_path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a document")

    with pytest.raises(AppException) as error:
        TextExtractionService().extract_text(path)

    assert error.value.code == "TEXT_EXTRACTION_FAILED"
