"""Extract selectable text from validated PDF and DOCX files."""

from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document

from app.core.exceptions import AppException


@dataclass(frozen=True)
class TextExtractionResult:
    text: str
    page_count: int
    character_count: int

    def as_dict(self) -> dict[str, str | int]:
        return {
            "text": self.text,
            "page_count": self.page_count,
            "character_count": self.character_count,
        }


class TextExtractionService:
    def extract_text(self, file_path: str | Path, file_type: str | None = None) -> TextExtractionResult:
        path = Path(file_path)
        document_type = file_type or path.suffix.lower().lstrip(".")

        try:
            if document_type == "pdf":
                result = self._extract_pdf(path)
            elif document_type == "docx":
                result = self._extract_docx(path)
            else:
                raise AppException(
                    "UNSUPPORTED_FILE_TYPE",
                    "Text extraction supports only PDF and DOCX files.",
                )
        except AppException:
            raise
        except Exception as error:
            raise AppException(
                "TEXT_EXTRACTION_FAILED",
                "The resume could not be read as a valid document.",
            ) from error

        if not result.text.strip():
            raise AppException(
                "NO_EXTRACTABLE_TEXT",
                "The resume does not contain extractable text.",
            )
        return result

    @staticmethod
    def _extract_pdf(path: Path) -> TextExtractionResult:
        pages: list[str] = []
        with fitz.open(path) as document:
            page_count = document.page_count
            for page in document:
                pages.append(page.get_text("text").strip())
        text = "\n\n".join(page for page in pages if page)
        return TextExtractionResult(text, page_count, len(text))

    @staticmethod
    def _extract_docx(path: Path) -> TextExtractionResult:
        document = Document(path)
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        text = "\n".join(blocks)
        return TextExtractionResult(text, 1, len(text))
